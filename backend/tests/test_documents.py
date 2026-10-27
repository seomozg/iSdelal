"""
Unit and API tests for document ingestion (PDF/XLSX/Google Docs/Sheets).
Run with: pytest backend/tests/test_documents.py -v
"""
import asyncio
import io
import uuid
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.auth import create_jwt
from app.documents import (
    MAX_FILE_SIZE,
    DocumentError,
    delete_document_chunks,
    google_export_target,
    index_text,
    parse_pdf,
    parse_xlsx,
    validate_upload,
)


def minimal_pdf(text: str) -> bytes:
    """Build a tiny single-page PDF containing `text` (no extra deps needed)."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


class TestValidation:
    """validate_upload: extension + size limits."""

    def test_pdf_ok(self):
        assert validate_upload("price-list.PDF", 1000) == "pdf"

    def test_xlsx_ok(self):
        assert validate_upload("report.xlsx", 1000) == "xlsx"

    def test_rejects_other_extension(self):
        with pytest.raises(DocumentError, match="Поддерживаются только файлы"):
            validate_upload("notes.txt", 1000)

    def test_rejects_too_big(self):
        with pytest.raises(DocumentError, match="больше 10 МБ"):
            validate_upload("big.pdf", MAX_FILE_SIZE + 1)

    def test_rejects_empty_name(self):
        with pytest.raises(DocumentError):
            validate_upload("", 100)


class TestParsing:
    """PDF and Excel text extraction."""

    def test_parse_pdf_extracts_text(self):
        text = parse_pdf(minimal_pdf("Hello world this pdf has plenty of words to extract"))
        assert "Hello world" in text

    def test_parse_pdf_rejects_garbage(self):
        with pytest.raises(DocumentError):
            parse_pdf(b"this is not a pdf at all")

    def test_parse_xlsx_reads_all_sheets(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Prices"
        ws.append(["Product", "Price"])
        ws.append(["Padel racket", 4500])
        ws2 = wb.create_sheet("Stock")
        ws2.append(["Warehouse", "Qty"])
        ws2.append(["Moscow", 12])
        buf = io.BytesIO()
        wb.save(buf)

        text = parse_xlsx(buf.getvalue())
        assert "Padel racket" in text
        assert "4500" in text
        assert "Moscow" in text

    def test_parse_xlsx_rejects_empty(self):
        wb = Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(DocumentError):
            parse_xlsx(buf.getvalue())


class TestGoogleUrls:
    """Google Docs/Sheets URL classification and fetching."""

    def test_doc_url(self):
        source_type, export_url = google_export_target(
            "https://docs.google.com/document/d/abc123XYZ/edit?usp=sharing"
        )
        assert source_type == "google_doc"
        assert export_url == "https://docs.google.com/document/d/abc123XYZ/export?format=txt"

    def test_sheet_url(self):
        source_type, export_url = google_export_target(
            "https://docs.google.com/spreadsheets/d/sheet42/edit#gid=0"
        )
        assert source_type == "google_sheet"
        assert export_url == "https://docs.google.com/spreadsheets/d/sheet42/export?format=csv"

    def test_rejects_foreign_url(self):
        with pytest.raises(DocumentError, match="Не распознана ссылка"):
            google_export_target("https://example.com/document/d/abc")

    def test_fetch_rejects_unpublished_doc(self, monkeypatch):
        import app.documents as documents_module

        resp = MagicMock()
        resp.status_code = 200
        resp.headers = {"Content-Type": "text/html; charset=utf-8"}
        monkeypatch.setattr(documents_module.requests, "get", lambda *a, **k: resp)

        with pytest.raises(DocumentError, match="Опубликуйте"):
            documents_module.fetch_google_document(
                "https://docs.google.com/document/d/abc/edit"
            )

    def test_fetch_sheet_converts_csv(self, monkeypatch):
        import app.documents as documents_module

        resp = MagicMock()
        resp.status_code = 200
        resp.headers = {"Content-Type": "text/csv"}
        resp.text = "Name,Price\nWidget,99\n"
        monkeypatch.setattr(documents_module.requests, "get", lambda *a, **k: resp)

        text = documents_module.fetch_google_document(
            "https://docs.google.com/spreadsheets/d/s1/edit"
        )
        assert "Widget | 99" in text

class TestQdrantIndexing:
    """index_text / delete_document_chunks against a mocked Qdrant client."""

    @pytest.fixture(autouse=True)
    def _mock_embeds(self, monkeypatch):
        monkeypatch.setattr(
            "app.documents.embed_texts",
            lambda texts: ([[0.1] * 8] * len(texts), 5),
        )

    def test_index_text_creates_collection_and_points(self, monkeypatch):
        q = MagicMock()
        q.get_collections.return_value = SimpleNamespace(collections=[])
        monkeypatch.setattr("app.documents.get_qdrant_client", lambda: q)

        n = index_text("word " * 120, "col_a", 55, "pdf", "a.pdf")

        q.create_collection.assert_called_once()
        kwargs = q.upsert.call_args.kwargs
        assert kwargs["collection_name"] == "col_a"
        assert len(kwargs["points"]) == n and n > 0
        payload = kwargs["points"][0]["payload"]
        assert payload["document_id"] == 55
        assert payload["source_type"] == "pdf"
        assert payload["url"] == "a.pdf"
        assert "text" in payload

    def test_index_text_reuses_existing_collection(self, monkeypatch):
        q = MagicMock()
        q.get_collections.return_value = SimpleNamespace(
            collections=[SimpleNamespace(name="col_b")]
        )
        monkeypatch.setattr("app.documents.get_qdrant_client", lambda: q)

        index_text("word " * 60, "col_b", 1, "google_doc", "doc1")

        q.create_collection.assert_not_called()
        q.upsert.assert_called_once()

    def test_index_text_raises_without_embeddings(self, monkeypatch):
        monkeypatch.setattr("app.documents.embed_texts", lambda texts: ([], 0))
        with pytest.raises(DocumentError, match="эмбеддинг"):
            index_text("word " * 60, "col_c", 1, "pdf", "a.pdf")

    def test_delete_document_chunks(self, monkeypatch):
        q = MagicMock()
        q.count.return_value = SimpleNamespace(count=4)
        monkeypatch.setattr("app.documents.get_qdrant_client", lambda: q)

        assert delete_document_chunks(7, "col_d") == 4

        selector = q.delete.call_args.kwargs["points_selector"]
        condition = selector.must[0]
        assert condition.key == "document_id"
        assert condition.match.value == 7

    def test_delete_without_collection_returns_zero(self, monkeypatch):
        q = MagicMock()
        q.count.side_effect = Exception("collection not found")
        monkeypatch.setattr("app.documents.get_qdrant_client", lambda: q)

        assert delete_document_chunks(1, "missing") == 0


# ------------------------------------------------------------------ API helpers

def _create_owner():
    """Create a fresh user with one site; returns (auth_headers, site_id)."""
    tag = uuid.uuid4().hex[:12]

    async def inner():
        from app.database import async_session_factory
        from app.models import Site, User

        async with async_session_factory() as s:
            user = User(email=f"doc_{tag}@test.dev", google_id=f"gid_{tag}")
            s.add(user)
            await s.flush()
            site = Site(user_id=user.id, collection_name=f"doc_col_{tag}", url="https://x.dev")
            s.add(site)
            await s.commit()
            return user.id, site.id

    user_id, site_id = asyncio.run(inner())
    token = create_jwt(user_id, f"doc_{tag}@test.dev")
    return {"Authorization": f"Bearer {token}"}, site_id


def _add_doc(site_id: int, name: str = "old.pdf") -> int:
    async def inner():
        from app.database import async_session_factory
        from app.models import Document

        async with async_session_factory() as s:
            doc = Document(site_id=site_id, name=name, source_type="pdf", status="completed")
            s.add(doc)
            await s.commit()
            return doc.id

    return asyncio.run(inner())


def _get_doc(doc_id: int):
    async def inner():
        from app.database import async_session_factory
        from app.models import Document

        async with async_session_factory() as s:
            return await s.get(Document, doc_id)

    return asyncio.run(inner())


@pytest.fixture
def client():
    from fastapi.testclient import TestClient
    from app.main import app

    return TestClient(app, raise_server_exceptions=False)


@pytest.fixture
def fake_background(monkeypatch):
    """Record background ingest calls instead of running them."""
    calls = []

    async def fake(document_id, file_bytes=None, google_url=None):
        calls.append((document_id, file_bytes, google_url))

    monkeypatch.setattr("app.routers.documents.ingest_document_background", fake)
    return calls


@pytest.fixture
def owner():
    return _create_owner()


class TestDocumentApi:
    """REST endpoints under /api/sites/{id}/documents."""

    def test_upload_pdf_accepted(self, client, owner, fake_background):
        headers, site_id = owner
        pdf = minimal_pdf("document words for upload test plenty of text")
        resp = client.post(
            f"/api/sites/{site_id}/documents",
            files={"file": ("price.pdf", pdf, "application/pdf")},
            headers=headers,
        )
        assert resp.status_code == 202, resp.text
        data = resp.json()
        assert data["status"] == "processing"
        assert data["source_type"] == "pdf"
        assert data["name"] == "price.pdf"
        assert fake_background and fake_background[0][1] == pdf

    def test_add_google_url(self, client, owner, fake_background):
        headers, site_id = owner
        resp = client.post(
            f"/api/sites/{site_id}/documents",
            data={"url": "https://docs.google.com/document/d/abc123/edit"},
            headers=headers,
        )
        assert resp.status_code == 202, resp.text
        assert resp.json()["source_type"] == "google_doc"
        assert fake_background[0][2].endswith("/abc123/edit")

    def test_rejects_unsupported_file(self, client, owner):
        headers, site_id = owner
        resp = client.post(
            f"/api/sites/{site_id}/documents",
            files={"file": ("notes.txt", b"hello", "text/plain")},
            headers=headers,
        )
        assert resp.status_code == 400
        assert "Поддерживаются только файлы" in resp.json()["detail"]

    def test_rejects_bad_url(self, client, owner):
        headers, site_id = owner
        resp = client.post(
            f"/api/sites/{site_id}/documents",
            data={"url": "https://example.com/not-a-google-doc"},
            headers=headers,
        )
        assert resp.status_code == 400

    def test_rejects_file_and_url_together(self, client, owner):
        headers, site_id = owner
        resp = client.post(
            f"/api/sites/{site_id}/documents",
            files={"file": ("a.pdf", minimal_pdf("words " * 10), "application/pdf")},
            data={"url": "https://docs.google.com/document/d/x/edit"},
            headers=headers,
        )
        assert resp.status_code == 400

    def test_requires_auth(self, client, owner):
        _, site_id = owner
        resp = client.get(f"/api/sites/{site_id}/documents")
        assert resp.status_code == 401

    def test_unknown_site_is_404(self, client, owner):
        headers, _ = owner
        assert client.get("/api/sites/999999/documents", headers=headers).status_code == 404
        assert client.post(
            "/api/sites/999999/documents",
            data={"url": "https://docs.google.com/document/d/x/edit"},
            headers=headers,
        ).status_code == 404

    def test_list_documents(self, client, owner):
        headers, site_id = owner
        _add_doc(site_id, name="contract.pdf")
        resp = client.get(f"/api/sites/{site_id}/documents", headers=headers)
        assert resp.status_code == 200
        body = resp.json()
        assert "contract.pdf" in [d["name"] for d in body]
        assert all(d["site_id"] == site_id for d in body)

    def test_delete_document(self, client, owner, monkeypatch):
        headers, site_id = owner
        doc_id = _add_doc(site_id, name="trash.pdf")
        monkeypatch.setattr(
            "app.routers.documents.delete_document_chunks", lambda d, c: 3
        )

        resp = client.delete(
            f"/api/sites/{site_id}/documents/{doc_id}", headers=headers
        )
        assert resp.status_code == 200
        assert resp.json() == {"status": "deleted", "chunks_removed": 3}
        assert _get_doc(doc_id) is None

    def test_delete_foreign_document_is_404(self, client, owner):
        headers, site_id = owner
        alien_doc = _add_doc(999999, name="alien.pdf")  # belongs to another site
        resp = client.delete(
            f"/api/sites/{site_id}/documents/{alien_doc}", headers=headers
        )
        assert resp.status_code == 404
        assert _get_doc(alien_doc) is not None

